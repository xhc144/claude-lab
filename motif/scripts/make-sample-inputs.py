#!/usr/bin/env python3
"""Generates deliberately plain source materials (default-template PPTX, DOCX, XLSX, Markdown,
PDF) used to test the importers and to demonstrate turning ordinary content into designed motion.
Requires python-pptx, python-docx, openpyxl, Pillow; LibreOffice (soffice) for the PDF."""
import os, subprocess, sys
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw

out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), '..', 'examples', 'inputs')
os.makedirs(out, exist_ok=True)

# a simple product photo stand-in
img = Image.new('RGB', (1200, 800), (232, 236, 242))
d = ImageDraw.Draw(img)
for i in range(40):
    d.ellipse((600 - 300 + i * 4, 400 - 220 + i * 3, 600 + 300 - i * 4, 400 + 220 - i * 3), outline=(80 + i * 3, 110 + i * 2, 200), width=2)
d.rounded_rectangle((420, 300, 780, 500), radius=40, fill=(30, 60, 140))
d.text((470, 380), 'MOTIF CLOUD', fill=(255, 255, 255))
img_path = os.path.join(out, 'product.png')
img.save(img_path)

# ---------------- PPTX: a plain default-template business deck ----------------
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = '2025 年度经营汇报'
s.placeholders[1].text = '星云科技 · 战略与运营部'
s.notes_slide.notes_text_frame.text = '各位好，下面汇报 2025 年的经营情况。'

s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = '年度亮点'
tf = s.placeholders[1].text_frame
tf.text = '营收达到 86 亿元，同比增长 65%'
for t in ['付费用户突破 380 万', '海外收入占比提升到 23%', '研发投入占营收 18%', '首次实现全年盈利']:
    tf.add_paragraph().text = t
s.notes_slide.notes_text_frame.text = '今年最重要的变化是规模和盈利同时突破。'

s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = '营收五年走势（亿元）'
cd = CategoryChartData()
cd.categories = ['2021', '2022', '2023', '2024', '2025']
cd.add_series('营收', (12, 19, 31, 52, 86))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5), cd)

s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = '收入结构'
cd = CategoryChartData()
cd.categories = ['订阅', '企业服务', '广告', '其他']
cd.add_series('占比', (54, 28, 12, 6))
s.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(2), Inches(1.5), Inches(6), Inches(5), cd)

s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = '区域表现'
rows, cols = 5, 4
tbl = s.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(8.4), Inches(3)).table
data = [['区域', '营收（亿）', '同比', '占比'], ['华东', '31.2', '+58%', '36%'], ['华南', '22.4', '+71%', '26%'], ['海外', '19.6', '+124%', '23%'], ['其他', '12.8', '+22%', '15%']]
for r in range(rows):
    for c in range(cols):
        tbl.cell(r, c).text = data[r][c]

s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = '发展历程'
tf = s.placeholders[1].text_frame
tf.text = '2019 公司成立'
for t in ['2021 完成 A 轮融资', '2023 用户突破百万', '2024 进入东南亚', '2025 实现盈利']:
    tf.add_paragraph().text = t

s = prs.slides.add_slide(prs.slide_layouts[8])
s.shapes.title.text = '新一代产品 Motif Cloud'
s.placeholders[1].insert_picture(img_path)
s.placeholders[2].text = '面向中小企业的一站式数据平台，十分钟完成接入。'

s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = '2026 计划'
tf = s.placeholders[1].text_frame
tf.text = '第一步：完善产品矩阵'
for t in ['第二步：深耕重点行业', '第三步：扩大海外布局', '第四步：建设开发者生态']:
    tf.add_paragraph().text = t

s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = '谢谢'
s.placeholders[1].text = 'strategy@example.com'
prs.save(os.path.join(out, 'annual-review.pptx'))

# ---------------- DOCX: a research note ----------------
doc = Document()
doc.add_heading('城市骑行调研报告', 0)
doc.add_paragraph('本报告基于 2025 年对 12 个城市、共 3,600 名骑行者的问卷与访谈。')
doc.add_heading('一、核心发现', 1)
doc.add_paragraph('通勤骑行比例达到 41%，比三年前提升 17 个百分点。平均单次骑行距离为 4.2 公里。')
doc.add_heading('骑行的主要原因', 2)
for t in ['省时：避开拥堵', '健康：日常锻炼', '省钱：替代打车', '环保：减少排放']:
    doc.add_paragraph(t, style='List Bullet')
doc.add_heading('二、城市对比', 1)
t = doc.add_table(rows=1, cols=3)
t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = '城市', '骑行通勤占比', '自行车道（公里）'
for row in [['深圳', '46%', '2,800'], ['成都', '44%', '2,300'], ['杭州', '41%', '2,100'], ['北京', '38%', '3,200'], ['上海', '35%', '2,600']]:
    cells = t.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = v
doc.add_heading('三、建议', 1)
for t2 in ['加密社区微循环车道', '地铁站增设停车点', '推广电动助力车标准']:
    doc.add_paragraph(t2, style='List Number')
doc.add_paragraph('骑行不是替代，而是城市交通的毛细血管。', style='Intense Quote')
doc.save(os.path.join(out, 'cycling-report.docx'))

# ---------------- XLSX: monthly metrics ----------------
wb = Workbook()
ws = wb.active
ws.title = '月度活跃'
ws.append(['月份', '2024', '2025'])
for m, a, b in zip(['1月', '2月', '3月', '4月', '5月', '6月', '7月', '8月', '9月', '10月', '11月', '12月'], [80, 84, 90, 95, 99, 104, 110, 118, 122, 130, 141, 150], [120, 132, 141, 160, 175, 190, 210, 236, 250, 281, 320, 380]):
    ws.append([m, a, b])
ws2 = wb.create_sheet('渠道占比')
ws2.append(['渠道', '占比'])
for r in [['自然搜索', 0.38], ['社交媒体', 0.27], ['合作伙伴', 0.21], ['广告', 0.14]]:
    ws2.append(r)
for row in ws2.iter_rows(min_row=2, min_col=2, max_col=2):
    row[0].number_format = '0%'
wb.save(os.path.join(out, 'metrics.xlsx'))

# ---------------- Markdown lesson ----------------
with open(os.path.join(out, 'pythagoras.md'), 'w', encoding='utf-8') as f:
    f.write('''# 勾股定理

两千五百年前，人们就发现了直角三角形的秘密。

## 定理内容

$$ a^2 + b^2 = c^2 $$

直角三角形两条直角边的平方和，等于斜边的平方。

## 证明思路

1. 画出边长为 a+b 的大正方形
2. 在内部摆放四个全等的直角三角形
3. 比较中间小正方形的面积
4. 得到 a² + b² = c²

## 生活中的应用

- 测量：不用梯子也能算出树高
- 建筑：用 3-4-5 绳结确定直角
- 导航：计算两点间的直线距离

> 数学是上帝用来书写宇宙的文字。 —— 伽利略
''')

# ---------------- PDF (via LibreOffice) ----------------
try:
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', out, os.path.join(out, 'cycling-report.docx')], check=True, capture_output=True, timeout=120)
except Exception as e:
    print('PDF conversion skipped:', e)
print('sample inputs written to', os.path.abspath(out))
